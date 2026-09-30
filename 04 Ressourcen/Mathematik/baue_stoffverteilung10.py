#!/usr/bin/env python3
"""Stoffverteilungsplan Mathe Klasse 10 (10b) 2026/27, übertragen aus dem Plan 25/26 (gleiche Themenfolge, gleiche Optik).
Termine: KA und Abschlussprüfung laut Infoblatt 10b, Ferien laut 04 Ressourcen/Physik/Rahmendaten Schuljahr 2026-27.md.
Aufruf: python3 baue_stoffverteilung10.py  -> HTML, PDF und XLSX 'Stoffverteilungsplan 26-27 Mathe 10' (im Ordner des Skripts)."""
import subprocess
from datetime import date, timedelta
from pathlib import Path

HIER = Path(__file__).resolve().parent
NAME = "Stoffverteilungsplan 26-27 Mathe 10"
CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"

# (Wochennummer, Montag, Thema, Termin, rot?)
W = [
    (1, date(2026, 9, 14), "Organisatorisches, Wdh Basiswissen, Zufall", "", False),
    (2, date(2026, 9, 21), "2 stufige Versuche / mit R ohne R", "", False),
    (3, date(2026, 9, 28), "Kurztest Wahrscheinlichkeit; Beginn Prozentrechnen (Wdh)", "Kurztest Mi 30.9.", False),
    (4, date(2026, 10, 5), "G+ / G- / Rabatt / Skonto", "", False),
    (5, date(2026, 10, 12), "Kpt / Zinseszins", "", False),
    (6, date(2026, 10, 19), "Zeitfenster", "", False),
    "Herbstferien 26.10.-30.10.",
    (7, date(2026, 11, 2), "KA 1 + Wachstumssparen", "KA Mo 2.11.", True),
    (8, date(2026, 11, 9), "Seitenverhältnisse im rechtwinkeligen Dreieck / sin, cos, tan", "", False),
    (9, date(2026, 11, 16), "Beliebige Dreiecke / Vielecke / Trigo im Raum", "", False),
    (10, date(2026, 11, 23), "Sinus am Einheitskreis", "", False),
    (11, date(2026, 11, 30), "6 seitige Pyramide V+O / n-seitige Pyramide", "", False),
    (12, date(2026, 12, 7), "Kegel V+O / Kugel V+O", "", False),
    (13, date(2026, 12, 14), "KA 2", "KA Fr 18.12.", True),
    (14, date(2026, 12, 21), "Zusammengesetzte Körper / Wdh Prisma", "nur Mo/Di", False),
    "Weihnachtsferien 23.12.-9.1.",
    (15, date(2027, 1, 11), "Streckenzüge auf Körpern", "", False),
    (16, date(2027, 1, 18), "Daten Boxplot", "", False),
    (17, date(2027, 1, 25), "Potenzen ganze Zahl als Exponent U gl. Basis", "", False),
    (18, date(2027, 2, 1), "Potenzen gl. Exp / wissenschaftl. Schreibweise, KA 3", "KA Do 4.2.", True),
    "Faschingsferien 8.2.-12.2. (Annahme, laut Schulkalender prüfen)",
    (19, date(2027, 2, 15), "LGS - BGL - quadr. Gleichungen WDH", "", False),
    (20, date(2027, 2, 22), "Geradengleichung y=mx+b / Punktprobe / Punkt-Steigungsformel", "", False),
    (21, date(2027, 3, 1), "Com Ex", "ComEx (Termin prüfen)", False),
    (22, date(2027, 3, 8), "y=x² / y=ax² / y=x²+c", "", False),
    (23, date(2027, 3, 15), "Normalform / Scheitelform", "", False),
    (24, date(2027, 3, 22), "Schnittpunkte / Nullstellen / Modellieren", "nur Mo-Mi", False),
    "Osterferien 25.3.-4.4. (Gründonnerstag bis Ostermontag frei)",
    (25, date(2027, 4, 5), "KA 4", "KA Di 6.4.", True),
    (26, date(2027, 4, 12), "Prüfungsvorbereitung", "", False),
    (27, date(2027, 4, 19), "Prüfungsvorbereitung", "", False),
    (28, date(2027, 4, 26), "Prüfungswoche", "Abschlussprüfung Mathe Do 29.4.", True),
    (29, date(2027, 5, 3), "Puffer", "Do Chr. Himmelfahrt", False),
    (30, date(2027, 5, 10), "Puffer", "", False),
    "Pfingstferien 17.5.-29.5. (Mo 17.5. Pfingstmontag)",
    (31, date(2027, 5, 31), "Puffer", "", False),
    (32, date(2027, 6, 7), "Puffer", "", False),
    (33, date(2027, 6, 14), "Puffer", "", False),
    (34, date(2027, 6, 21), "Puffer", "", False),
    (35, date(2027, 6, 28), "Puffer", "", False),
    (36, date(2027, 7, 5), "Puffer", "", False),
    (37, date(2027, 7, 12), "Puffer", "", False),
    (38, date(2027, 7, 19), "Puffer", "", False),
    (39, date(2027, 7, 26), "", "letzter Schultag Mi 28.7.", False),
]
FUSS = ("Offen: Termine der mündlichen Prüfung und der Entlassfeier der 10er (Vorjahr: mündliche Prüfung in Woche 35, "
        "Entlassung freitags in Woche 36). Die Faschingsferien sind eine Annahme.")


def d(x):
    return f"{x.day}.{x.month}"


def zeilen():
    for r in W:
        if isinstance(r, str):
            yield ("ferien", r)
        else:
            nr, mo, thema, termin, rot = r
            yield ("woche", nr, d(mo), d(mo + timedelta(days=4)), thema, termin, rot)


def html():
    rows = []
    for z in zeilen():
        if z[0] == "ferien":
            rows.append(f'<tr class="fe"><td colspan="5">{z[1]}</td></tr>')
        else:
            _, nr, von, bis, thema, termin, rot = z
            c = ' class="rot"' if rot else ""
            rows.append(f'<tr><td class="n">{nr}</td><td class="dt">{von}</td><td class="dt">{bis}</td>'
                        f'<td class="th"><span{c}>{thema}</span></td><td class="tm"><span{c}>{termin}</span></td></tr>')
    return f"""<!DOCTYPE html><html lang="de"><head><meta charset="utf-8"><title>{NAME}</title><style>
@page{{size:A4;margin:9mm 10mm}}
*{{box-sizing:border-box;-webkit-print-color-adjust:exact;print-color-adjust:exact}}
body{{margin:0;font:8.6pt/1.2 Arial,Helvetica,sans-serif;color:#000}}
.kopf{{background:#ffff00;border:.3mm solid #000;text-align:center;font-weight:700;padding:1.2mm 0;font-size:9.5pt}}
table{{border-collapse:collapse;width:100%}}
td{{border:.2mm solid #000;padding:.9mm 1.6mm}}
tr.h td{{font-weight:700;border-bottom:.3mm solid #000}}
.n{{width:7mm;text-align:right}}.dt{{width:11mm;text-align:right}}.tm{{width:39mm;font-size:7.4pt}}
.fe td{{text-align:center;font-size:8pt}}
.rot{{color:#e00000}}
.fuss{{margin-top:2mm;font-size:7.2pt;color:#444}}
</style></head><body>
<div class="kopf">Stoffverteilungsplan 26/27 &nbsp;&nbsp;&nbsp; Klasse 10 &nbsp;&nbsp;&nbsp; Fach Mathe &nbsp;&nbsp; Lehrer Klein</div>
<table><tr class="h"><td class="n">Wo</td><td colspan="2" style="text-align:center">Datum</td><td>Thema</td><td>Termine</td></tr>
{"".join(rows)}</table>
<div class="fuss">{FUSS}</div></body></html>"""


def xlsx(pfad):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    wb = Workbook()
    ws = wb.active
    ws.title = "Klasse 10"
    dünn = Side(style="thin")
    rand = Border(left=dünn, right=dünn, top=dünn, bottom=dünn)
    ws.merge_cells("A1:F1")
    ws["A1"] = "Stoffverteilungsplan 26/27      Klasse 10      Fach Mathe      Lehrer Klein"
    ws["A1"].font = Font(bold=True)
    ws["A1"].fill = PatternFill("solid", fgColor="FFFF00")
    ws["A1"].alignment = Alignment(horizontal="center")
    for col, t in enumerate(["Wo", "von", "bis", "Thema", "Termine"], start=1):
        c = ws.cell(row=2, column=col if col < 4 else col, value=t)
        c.font = Font(bold=True)
        c.border = rand
    r = 3
    for z in zeilen():
        if z[0] == "ferien":
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
            c = ws.cell(row=r, column=1, value=z[1])
            c.alignment = Alignment(horizontal="center")
        else:
            _, nr, von, bis, thema, termin, rot = z
            for col, v in enumerate([nr, von, bis, thema, termin], start=1):
                c = ws.cell(row=r, column=col, value=v)
                c.border = rand
                if rot and col >= 4:
                    c.font = Font(color="E00000")
        r += 1
    ws.cell(row=r + 1, column=1, value=FUSS).font = Font(italic=True, size=9)
    for col, w in zip("ABCDE", (5, 8, 8, 62, 32)):
        ws.column_dimensions[col].width = w
    wb.save(pfad)


def main():
    h = HIER / f"{NAME}.html"
    h.write_text(html(), encoding="utf-8")
    subprocess.run([CHROME, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                    f"--print-to-pdf={HIER / (NAME + '.pdf')}", h.as_uri()], check=True, capture_output=True)
    xlsx(HIER / f"{NAME}.xlsx")
    print("geschrieben:", NAME, "(.html/.pdf/.xlsx)")


if __name__ == "__main__":
    main()
